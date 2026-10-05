#!/usr/bin/env python3
"""
download_logos.py
Part of the WeMap data pipeline.

Concurrently downloads all company logo images from Google Drive links in the WeMap
dataset and caches them locally in a logos/ directory for fast, offline, and reliable serving.

Usage:
    python download_logos.py                          # auto-detects *_geocoded.xlsx in CWD
    python download_logos.py path/to/geocoded.xlsx
    python download_logos.py path/to/geocoded.xlsx --logos-dir /path/to/logos
    python download_logos.py path/to/geocoded.xlsx --workers 20

RUN THIS LOCALLY, NOT IN A SANDBOXED/NETWORK-RESTRICTED ENVIRONMENT.
"""

import re
import sys
import time
import urllib.request
import argparse
import openpyxl
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path


def get_extension_from_headers_and_data(content_type, data):
    if "svg" in content_type or data.lstrip().startswith(b"<svg") or b"<svg" in data[:100]:
        return ".svg"
    elif "jpeg" in content_type or "jpg" in content_type or data.startswith(b"\xff\xd8\xff"):
        return ".jpg"
    elif "gif" in content_type or data.startswith(b"GIF8"):
        return ".gif"
    elif "webp" in content_type or b"WEBP" in data[:20]:
        return ".webp"
    elif "png" in content_type or data.startswith(b"\x89PNG"):
        return ".png"
    return ".png"


def download_logo(org_id, url, logos_dir):
    """Download a single logo; returns (org_id, rel_path_or_None, status_str)."""
    if not url or not str(url).strip().startswith("http"):
        return org_id, None, "No URL"

    url = str(url).strip()

    # Skip if already downloaded (any extension)
    existing = list(logos_dir.glob(f"{org_id}.*"))
    if existing:
        return org_id, f"logos/{existing[0].name}", "Cached"

    # Convert Google Drive share links to direct-download URLs
    m = re.search(r"id=([a-zA-Z0-9_-]+)", url)
    if "drive.google.com" in url and m:
        direct_url = f"https://drive.google.com/uc?export=view&id={m.group(1)}"
    else:
        direct_url = url

    headers = {
        "User-Agent": (
            "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
            "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
        )
    }

    for _ in range(2):
        try:
            req = urllib.request.Request(direct_url, headers=headers)
            with urllib.request.urlopen(req, timeout=10) as response:
                data = response.read()
                # Reject HTML error pages
                if len(data) < 100 or b"<!DOCTYPE html>" in data[:200] or b"<html" in data[:200]:
                    time.sleep(1)
                    continue
                content_type = response.headers.get("Content-Type", "").lower()
                ext = get_extension_from_headers_and_data(content_type, data)
                filename = f"{org_id}{ext}"
                (logos_dir / filename).write_bytes(data)
                return org_id, f"logos/{filename}", "Downloaded"
        except Exception:
            time.sleep(1)

    return org_id, None, "Failed"


def main():
    parser = argparse.ArgumentParser(
        description="Download org logos from a WeMap dataset into a local logos/ directory."
    )
    parser.add_argument(
        "xlsx_path",
        nargs="?",
        default=None,
        help=(
            "Path to geocoded xlsx file. "
            "If omitted, auto-detects *_geocoded.xlsx then any *.xlsx in the current directory."
        ),
    )
    parser.add_argument(
        "--logos-dir",
        default=None,
        help="Directory to save logos (default: logos/ next to the xlsx file)",
    )
    parser.add_argument(
        "--workers",
        type=int,
        default=15,
        help="Number of parallel download workers (default: 15)",
    )
    args = parser.parse_args()

    # --- Resolve xlsx path ---
    if args.xlsx_path:
        xlsx_path = Path(args.xlsx_path)
    else:
        candidates = sorted(Path(".").glob("*_geocoded.xlsx")) or sorted(Path(".").glob("*.xlsx"))
        if not candidates:
            print("Error: no xlsx file found. Pass a path as the first argument.")
            sys.exit(1)
        xlsx_path = candidates[0]
        print(f"Auto-detected dataset: {xlsx_path}")

    if not xlsx_path.exists():
        print(f"Error: file not found: {xlsx_path}")
        sys.exit(1)

    # --- Resolve logos directory ---
    logos_dir = Path(args.logos_dir) if args.logos_dir else xlsx_path.parent / "logos"
    logos_dir.mkdir(parents=True, exist_ok=True)
    print(f"Saving logos to: {logos_dir.resolve()}")

    # --- Load workbook ---
    wb = openpyxl.load_workbook(xlsx_path)
    ws = wb.active
    headers = [cell.value for cell in ws[1]]

    required = {"org_id", "WeMap_URL"}
    missing = required - set(headers)
    if missing:
        print(f"Error: missing required columns {missing} in {xlsx_path}.")
        print(f"  Found columns: {headers[:15]}")
        sys.exit(1)

    idx_id = headers.index("org_id")
    idx_url = headers.index("WeMap_URL")

    tasks = [
        (row[idx_id], row[idx_url])
        for row in ws.iter_rows(min_row=2, values_only=True)
        if row[idx_id] and row[idx_url] and str(row[idx_url]).strip().startswith("http")
    ]

    print(f"Total valid logo URLs in dataset: {len(tasks)}")
    print(f"Already downloaded logos: {len(list(logos_dir.glob('*.*')))}")

    downloaded = cached = failed = 0

    with ThreadPoolExecutor(max_workers=args.workers) as executor:
        futures = {executor.submit(download_logo, org_id, url, logos_dir): org_id for org_id, url in tasks}
        for count, future in enumerate(as_completed(futures), 1):
            _, rel_path, status = future.result()
            if rel_path:
                if status == "Downloaded":
                    downloaded += 1
                else:
                    cached += 1
            else:
                failed += 1
            if count % 100 == 0 or count == len(tasks):
                print(
                    f"Processed {count}/{len(tasks)} — "
                    f"Downloaded: {downloaded}, Cached: {cached}, Failed: {failed}"
                )

    print("Finished logo caching!")


if __name__ == "__main__":
    main()
